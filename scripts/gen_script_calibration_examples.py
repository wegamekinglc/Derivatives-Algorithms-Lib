#!/usr/bin/env python3
"""Generate direct-script yield and XCCY calibration workbooks."""

import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2022, 1, 1)
KNOTS = [
    date(2022, 1, 1),
    date(2022, 3, 15),
    date(2022, 6, 15),
    date(2022, 9, 21),
    date(2022, 12, 21),
    date(2023, 3, 15),
    date(2023, 6, 21),
    date(2023, 9, 20),
    date(2023, 12, 20),
    date(2024, 3, 15),
    date(2025, 1, 1),
    date(2027, 1, 1),
    date(2029, 1, 1),
    date(2032, 1, 1),
]
SWAPS = [
    (date(2022, 1, 1), date(2022, 1, 3), 1.00),
    (date(2022, 3, 15), date(2022, 6, 15), 1.05),
    (date(2022, 6, 15), date(2022, 9, 21), 1.12),
    (date(2022, 9, 21), date(2022, 12, 21), 1.16),
    (date(2022, 12, 21), date(2023, 3, 15), 1.21),
    (date(2023, 3, 15), date(2023, 6, 21), 1.27),
    (date(2023, 6, 21), date(2023, 9, 20), 1.45),
    (date(2023, 9, 20), date(2023, 12, 20), 1.68),
    (date(2023, 12, 20), date(2024, 3, 15), 1.92),
    (date(2022, 1, 1), date(2025, 1, 1), 1.68),
    (date(2022, 1, 1), date(2027, 1, 1), 2.10),
    (date(2022, 1, 1), date(2029, 1, 1), 2.20),
    (date(2022, 1, 1), date(2032, 1, 1), 2.07),
]
TITLE_FONT = Font(name="Calibri", size=13, bold=True)
HEADER_FONT = Font(name="Calibri", size=11, bold=True)
BODY_FONT = Font(name="Calibri", size=11)
MONO_FONT = Font(name="Consolas", size=10)
GRAY_FILL = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
THIN_BORDER = Border(*(Side(style="thin", color="D0D0D0"),) * 4)


def number_text(value):
    text = repr(float(value))
    return text[:-2] if text.endswith(".0") else text


def annual_periods(start, end):
    periods = []
    current = start
    while current < end:
        following = current.replace(year=current.year + 1)
        periods.append((current, min(following, end)))
        current = following
    return periods


def annual_swap_rows(start, end, coupon):
    rows = [("Definition", "N", "100"), ("Definition", "K", number_text(coupon))]
    periods = annual_periods(start, end)
    for accrual_start, accrual_end in periods:
        alpha = number_text((accrual_end - accrual_start).days / 365)
        rows.append(("Event", accrual_end.isoformat(), f"PAY PAYS N*({alpha}*K)"))
    for accrual_start, accrual_end in periods:
        alpha = number_text((accrual_end - accrual_start).days / 365)
        forward = (
            f"((1/FIX(IR[USD,DF,{accrual_end.isoformat()},{accrual_start.isoformat()}],"
            f"{accrual_start.isoformat()})-1)/{alpha})"
        )
        rows.append(
            ("Event", accrual_end.isoformat(), f"PAY PAYS N*(-{alpha}*{forward})")
        )
    return rows


def yield_instruments():
    result = []
    for index, (start, end, percent) in enumerate(SWAPS, 1):
        coupon = percent / 100
        result.append(
            {
                "id": f"ptirds-{index}",
                "start": start,
                "end": end,
                "quote": coupon,
                "quote_id": f"ptirds-{index}-quote",
                "observation_id": f"ptirds-{index}-observation",
                "measure": "ParRate",
                "measure_kind": "ParAffineRate",
                "unit": "Rate",
                "annuity_sign": 1,
                "notional": 100,
                "rows": annual_swap_rows(start, end, coupon),
            }
        )
    return result


def xccy_instrument():
    return {
        "id": "eur-usd",
        "start": date(2027, 5, 4),
        "end": date(2027, 11, 4),
        "quote": -0.002983490837363863,
        "quote_id": "eur-usd-basis",
        "observation_id": "fit-basis",
        "measure": "BasisSpread",
        "measure_kind": "ParAffineSpread",
        "unit": "BasisSpread",
        "annuity_sign": -1,
        "notional": 110,
        "rows": [
            ("Definition", "N", "110"),
            ("Definition", "K", "0.001"),
            (
                "Event",
                "2027-08-04",
                "PAY PAYS N*(1*1*0.25555555555555554*(FIX(IR[USD,LIBOR_3M_LCH],2027-04-30)+0)*1)",
            ),
            (
                "Event",
                "2027-11-04",
                "PAY PAYS N*(1*1*0.25555555555555554*(FIX(IR[USD,LIBOR_3M_LCH],2027-08-02)+0)*1)",
            ),
            (
                "Event",
                "2027-11-08",
                "PAY PAYS N*(-1*0.9090909090909091*0.5111111111111111*(FIX(IR[EUR,LIBOR_6M_LCH],2027-04-30)+K)*SPOT())",
            ),
            ("Event", "2027-05-04", "PAY PAYS N*(-1*1*1)"),
            ("Event", "2027-05-04", "PAY PAYS N*(1*0.9090909090909091*SPOT())"),
            ("Event", "2027-11-04", "PAY PAYS N*(1*1*1)"),
            ("Event", "2027-11-04", "PAY PAYS N*(-1*0.9090909090909091*SPOT())"),
        ],
    }


def write_table(ws, row, col, headers, rows):
    for offset, header in enumerate(headers):
        cell = ws.cell(row, col + offset, header)
        cell.font = HEADER_FONT
        cell.fill = GRAY_FILL
        cell.border = THIN_BORDER
    for i, values in enumerate(rows, 1):
        for j, value in enumerate(values):
            cell = ws.cell(row + i, col + j, value)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            if isinstance(value, date):
                cell.number_format = "YYYY-MM-DD"
    last_row = row + len(rows)
    last_col = col + len(headers) - 1
    return f"${get_column_letter(col)}${row}:${get_column_letter(last_col)}${last_row}"


def cell_ref(ws, row, col):
    title = f"'{ws.title}'" if " " in ws.title else ws.title
    return f"{title}!${get_column_letter(col)}${row}"


def write_definition(ws, row, definition):
    ws.cell(row, 1, f"{definition['id']} script definition").font = TITLE_FONT
    source = write_table(ws, row + 1, 1, ["kind", "key", "text"], definition["rows"])
    header = write_table(
        ws,
        row + 1,
        5,
        ["field", "value"],
        [
            ("pay_currency", "USD"),
            ("result_currency", "USD"),
            ("collateral_currency", "USD"),
            ("collateral_type", "OIS"),
            ("receiver", "PAY"),
            ("notional_parameter", "N"),
            ("settlement", "AtEvent"),
            ("valuation_cashflows", "IncludeUnsettledAtValuation"),
            ("discount_route", "USD.OIS"),
            ("convention_set", "dal.fixed-dates/1"),
        ],
    )
    unit = "Price:USD"
    parameters = write_table(
        ws,
        row + 1,
        8,
        [
            "key",
            "type",
            "unit",
            "role",
            "bindable",
            "lower",
            "upper",
            "lower_inclusive",
            "upper_inclusive",
        ],
        [
            ("N", "Real", unit, "Notional", True, 0, None, False, True),
            ("K", "Real", definition["unit"], "Coupon", True, None, None, True, True),
        ],
    )
    measures = write_table(
        ws,
        row + 1,
        19,
        [
            "id",
            "kind",
            "receiver",
            "parameter",
            "unit",
            "annuity_sign",
            "min_abs_annuity",
        ],
        [
            ("Price", "Price", "PAY", None, unit, None, None),
            (
                definition["measure"],
                definition["measure_kind"],
                "PAY",
                "K",
                definition["unit"],
                definition["annuity_sign"],
                1e-12,
            ),
        ],
    )
    handles = {}
    labels = ["definition", "instrument", "quote", "observation"]
    for offset, label in enumerate(labels):
        ws.cell(row + 1 + offset, 26, label.capitalize()).font = HEADER_FONT
    ws.cell(
        row + 1,
        27,
        f"=SCRIPTINSTRUMENTDEFINITION.NEW({source},{header},{parameters},{measures},)",
    ).font = MONO_FONT
    ws.cell(
        row + 2,
        27,
        f'=SCRIPTCALIBRATIONINSTRUMENT.NEW("{definition["id"]}",AA{row + 1},,'
        f"DATE({definition['end'].year},{definition['end'].month},{definition['end'].day}))",
    ).font = MONO_FONT
    ws.cell(
        row + 3,
        27,
        f'=CALIBRATIONMARKETQUOTE.NEW("{definition["quote_id"]}",'
        f'{number_text(definition["quote"])},"{definition["unit"]}",)',
    ).font = MONO_FONT
    ws.cell(
        row + 4,
        27,
        f'=CALIBRATIONOBSERVATION.NEW("{definition["observation_id"]}",AA{row + 2},'
        f'"{definition["measure"]}",AA{row + 3},1E-10)',
    ).font = MONO_FONT
    for offset, label in enumerate(labels):
        handles[label] = cell_ref(ws, row + 1 + offset, 27)
    return handles


def factory_layout(ws, label_col=1, value_col=2, field_col=4):
    handles = {}
    cursor = [4]

    def direct(label, formula):
        row = 4 + len(handles)
        ws.cell(row, label_col, label).font = HEADER_FONT
        ws.cell(row, value_col, formula).font = MONO_FONT
        handles[label] = f"${get_column_letter(value_col)}${row}"
        return handles[label]

    def fields(label, function, rows, suffix="", prefix=""):
        source = write_table(ws, cursor[0], field_col, ["field", "value"], rows)
        cursor[0] += len(rows) + 3
        return direct(label, f"={function}({prefix}{source}{suffix})")

    def table(label, function, headers, rows):
        source = write_table(ws, cursor[0], field_col, headers, rows)
        cursor[0] += len(rows) + 3
        return direct(label, f"={function}({source})")

    return handles, direct, fields, table


def add_notes(wb, title, lines):
    ws = wb.create_sheet("Notes")
    ws.cell(1, 1, title).font = TITLE_FONT
    ws.cell(3, 1, "Initialize DAL").font = HEADER_FONT
    ws.cell(3, 2, "=INIT.GLOBALDATA(1)").font = MONO_FONT
    for row, line in enumerate(lines, 5):
        ws.cell(row, 1, line).font = BODY_FONT
        ws.cell(row, 1).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 120
    ws.column_dimensions["B"].width = 28


def add_curve_scheme(wb, scheme):
    ws = wb.create_sheet(scheme)
    ws.cell(1, 1, f"Script yield calibration — {scheme}").font = TITLE_FONT
    ws.cell(
        2, 1, "13 scripted instruments, 14 knots, EXACT solve, AAD Jacobian"
    ).font = BODY_FONT
    handles, direct, fields, table = factory_layout(ws)
    write_table(ws, 4, 7, ["index", "knot_date"], list(enumerate(KNOTS, 1)))
    knot_values = f"$H$5:$H${4 + len(KNOTS)}"
    knot_handle = direct("Knot dates", f"=CALIBRATIONDATES.NEW({knot_values})")
    fields(
        "Geometry",
        "CURVEINSTRUMENTGEOMETRY.NEW",
        [
            ("today", TODAY),
            ("currency", "USD"),
            ("curve_name", "ptirds"),
            ("knot_dates", knot_handle),
            ("parameterization", "LOG_DISCOUNT"),
            ("knot_policy", "INPUT"),
            ("log_df_scheme", scheme),
            ("calibrate_discount_curve", True),
            ("target_collateral", "OIS"),
            ("libor_basis", "ACT_365F"),
        ],
    )
    route = table(
        "Routes",
        "CURVESCRIPTROUTES.NEW",
        ["id", "projection", "collateral", "tenor"],
        [["USD.OIS", False, "OIS", None]],
    )
    fields("Market", "CURVEINSTRUMENTMARKET.NEW", [("routes", route)])
    fixings = direct("Empty fixings", "=MARKETFIXINGSNAPSHOT.NEW(,,)")
    fields(
        "Valuation",
        "SCRIPTVALUATIONSETTINGS.NEW",
        [("evaluation_date", TODAY)],
        f",{fixings}",
        prefix='"ptirds-valuation",',
    )
    fields(
        "Solver",
        "CURVEINSTRUMENTSOLVER.NEW",
        [
            ("solve_mode", "EXACT"),
            ("initial_guess", 0.05),
            ("smoothing_weight", 1.0),
            ("max_evaluations", 200),
            ("max_restarts", 20),
            ("fit_norm_tolerance", 1e-8),
        ],
    )
    fields(
        "Options",
        "CURVEINSTRUMENTCALIBRATIONOPTIONS.NEW",
        [
            ("evaluation", "AutoFast"),
            ("jacobian_mode", "AAD"),
            ("difference_step", 1e-6),
            ("compute_quote_jacobian", False),
        ],
    )
    observations = direct(
        "Observations", "=CALIBRATIONOBSERVATIONS.NEW(Instruments!$E$4:$E$16)"
    )
    fields(
        "Spec",
        "CURVEINSTRUMENTCALIBRATIONSPEC.NEW",
        [
            ("geometry", handles["Geometry"]),
            ("market", handles["Market"]),
            ("valuation", handles["Valuation"]),
            ("observations", observations),
            ("solver", handles["Solver"]),
        ],
    )
    outcome = direct(
        "Calibration",
        f"=CALIBRATE.CURVEINSTRUMENTS({handles['Spec']},{handles['Options']})",
    )
    direct("Status", f"=CALIBRATIONOUTCOME.GET.STATUS({outcome})")
    report = direct("Report", f"=CALIBRATIONOUTCOME.GET.REPORT({outcome})")
    direct("Curve", f"=CALIBRATIONOUTCOME.GET.CURVE({outcome})")
    ws.cell(22, 1, "Observation report").font = TITLE_FONT
    ws.cell(23, 1, f"=CALIBRATIONREPORT.GET.OBSERVATIONS({report})").font = MONO_FONT
    ws.cell(40, 1, "Fitted parameters").font = TITLE_FONT
    ws.cell(41, 1, f"=CALIBRATIONREPORT.GET.PARAMETERS({report})").font = MONO_FONT
    ws.freeze_panes = "A4"


def generate_yield_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Instruments"
    ws.cell(
        1, 1, "Direct-script instruments for discount-curve interpolation"
    ).font = TITLE_FONT
    write_table(
        ws,
        3,
        1,
        ["instrument", "definition", "handle", "market_quote", "observation"],
        [],
    )
    definitions = yield_instruments()
    handles = []
    next_row = 20
    for definition in definitions:
        handles.append(write_definition(ws, next_row, definition))
        next_row += max(len(definition["rows"]), 10) + 5
    for i, definition in enumerate(definitions, 4):
        ws.cell(i, 1, definition["id"]).font = BODY_FONT
        for col, key in enumerate(
            ["definition", "instrument", "quote", "observation"], 2
        ):
            ws.cell(i, col, f"={handles[i - 4][key]}").font = MONO_FONT
    for scheme in ["LOG_LINEAR", "LOG_CUBIC_NATURAL", "MIXED"]:
        add_curve_scheme(wb, scheme)
    add_notes(
        wb,
        "012.interpolate_curve_script",
        [
            "The Instruments sheet constructs every definition from dated DAL script source rows. No native swap or native-to-script import function is used.",
            "The three calibration sheets share the same 13 instruments, quotes and 14 knots. Only log_discount_scheme changes: LOG_LINEAR, LOG_CUBIC_NATURAL and MIXED.",
            "ACT/365F accruals are unadjusted; the fixing date is the accrual start. N=100 and each K equals the original instrument's decimal par quote.",
            "The solver uses EXACT, initial guess 0.05, smoothing weight 1, 200 evaluations, 20 restarts and fit_norm_tolerance 1e-8. AAD is selected explicitly.",
            "Run INIT.GLOBALDATA(1), then recalculate the workbook. CALIBRATIONOUTCOME.GET.STATUS should return Succeeded on all three sheets.",
            "Observation and parameter tables are spilled from the typed report. This file contains formulas, not cached results; generating it does not execute Excel.",
        ],
    )
    for sheet in wb.worksheets:
        for cells in sheet.columns:
            letter = get_column_letter(cells[0].column)
            width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 115)
            sheet.column_dimensions[letter].width = max(12, width)
        sheet.sheet_view.zoomScale = 85
    ws.column_dimensions["C"].width = 115
    return wb


def generate_xccy_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Script Instrument"
    definition = xccy_instrument()
    source = write_definition(ws, 1, definition)
    ws = wb.create_sheet("Calibration")
    ws.cell(1, 1, "Direct-script XCCY basis calibration").font = TITLE_FONT
    ws.cell(
        2, 1, "USD/EUR fixed-notional swap; foreign spread; domestic collateral"
    ).font = BODY_FONT
    today, end = date(2026, 10, 21), date(2027, 11, 4)
    write_table(
        ws,
        4,
        1,
        ["curve_date", "USD_right_forward", "EUR_right_forward"],
        [(today, 0.04, 0.02), (end, 0.04, 0.02)],
    )
    handles, direct, fields, table = factory_layout(ws, 5, 6, 8)
    usd_curve = direct(
        "USD curve", '=DISCOUNTPWC.NEW("USD","USD",$A$5:$A$6,$B$5:$B$6,)'
    )
    eur_curve = direct(
        "EUR curve", '=DISCOUNTPWC.NEW("EUR","EUR",$A$5:$A$6,$C$5:$C$6,)'
    )
    usd_discounts = table(
        "USD discounts",
        "CALIBRATIONDISCOUNTCURVES.NEW",
        ["key", "curve"],
        [["OIS", usd_curve]],
    )
    usd_forwards = table(
        "USD forwards",
        "CALIBRATIONFORWARDCURVES.NEW",
        ["key", "curve"],
        [["QUARTERLY", usd_curve]],
    )
    usd_block = direct(
        "USD block",
        f'=CURVEBLOCK.NEW("USD","USD",{usd_discounts},{usd_forwards},"ACT_365F")',
    )
    eur_discounts = table(
        "EUR discounts",
        "CALIBRATIONDISCOUNTCURVES.NEW",
        ["key", "curve"],
        [["OIS", eur_curve]],
    )
    eur_forwards = table(
        "EUR forwards",
        "CALIBRATIONFORWARDCURVES.NEW",
        ["key", "curve"],
        [["SEMIANNUAL", eur_curve]],
    )
    eur_block = direct(
        "EUR block",
        f'=CURVEBLOCK.NEW("EUR","EUR",{eur_discounts},{eur_forwards},"ACT_365F")',
    )
    pair = direct("Currency pair", '=CURRENCYPAIR.NEW("USD","EUR")')
    fixings = direct("Empty fixings", "=MARKETFIXINGSNAPSHOT.NEW(,,)")
    fields(
        "Context",
        "XCCYSCRIPTCONTEXT.NEW",
        [
            ("valuation_time", today),
            ("pair", pair),
            ("collateral_currency", "USD"),
            ("discount_route", "USD.OIS"),
            ("fixings", fixings),
        ],
    )
    knots = direct("Basis knots", "=CALIBRATIONDATES.NEW($A$6)")
    observations = direct(
        "Observations", f"=CALIBRATIONOBSERVATIONS.NEW({source['observation']})"
    )
    fields(
        "Options",
        "CURVEINSTRUMENTCALIBRATIONOPTIONS.NEW",
        [
            ("evaluation", "AutoFast"),
            ("jacobian_mode", "AAD"),
            ("difference_step", 1e-6),
            ("compute_quote_jacobian", True),
        ],
    )
    fields(
        "Spec",
        "XCCYINSTRUMENTCALIBRATIONSPEC.NEW",
        [
            ("context", handles["Context"]),
            ("domestic_curve_block", usd_block),
            ("foreign_curve_block", eur_block),
            ("fx_spot", 1.1),
            ("knot_dates", knots),
            ("observations", observations),
        ],
    )
    outcome = direct(
        "Calibration",
        f"=CALIBRATE.XCCYINSTRUMENTS({handles['Spec']},{handles['Options']})",
    )
    direct("Status", f"=CALIBRATIONOUTCOME.GET.STATUS({outcome})")
    report = direct("Report", f"=CALIBRATIONOUTCOME.GET.REPORT({outcome})")
    direct("Basis curve", f"=CALIBRATIONOUTCOME.GET.CURVE({outcome})")
    ws.cell(27, 1, "Observation report").font = TITLE_FONT
    ws.cell(28, 1, f"=CALIBRATIONREPORT.GET.OBSERVATIONS({report})").font = MONO_FONT
    ws.cell(32, 1, "Fitted basis forward").font = TITLE_FONT
    ws.cell(33, 1, f"=CALIBRATIONREPORT.GET.PARAMETERS({report})").font = MONO_FONT
    ws.freeze_panes = "A4"
    add_notes(
        wb,
        "013.xccy_script_calibration",
        [
            "The Script Instrument sheet contains the complete direct source definition: two USD 3M coupons, one EUR 6M spread coupon, and initial/final principal exchanges with FX conversion.",
            "N=110 USD, foreign/domestic notional ratio=100/110, K=0.001, spread on the foreign leg, explicit USD.OIS collateral, and signed BasisSpread measure with annuity sign -1 and floor 1e-12.",
            "The calibration uses USD 4% and EUR 2% piecewise-constant forward curves, FX spot 1.1, one basis knot at 2027-11-04, and the original quote -0.002983490837363863.",
            "The default XCCY solver initial guess is 0. The basis parameterization is PIECEWISE_CONSTANT_FWD. Options select AAD and the quote Jacobian.",
            "Run INIT.GLOBALDATA(1), then recalculate. The expected basis forward is 0.003 and the outcome status is Succeeded.",
            "This file contains formulas, not cached results; generating it does not execute Excel.",
        ],
    )
    for sheet in wb.worksheets:
        for cells in sheet.columns:
            letter = get_column_letter(cells[0].column)
            width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 115)
            sheet.column_dimensions[letter].width = max(12, width)
        sheet.sheet_view.zoomScale = 85
    wb["Script Instrument"].column_dimensions["C"].width = 115
    return wb


def validate_examples():
    import dal

    def dal_date(value):
        return dal.Date_(value.year, value.month, value.day)

    def unit(token):
        if token.startswith("Price:"):
            return dal.CalibrationValueUnit_(
                kind="Price", currency=token.split(":", 1)[1]
            )
        return dal.CalibrationValueUnit_(kind=token)

    def instrument(definition):
        spec = dal.ScriptInstrumentDefinitionSpec_(
            source=dal.ScriptInstrumentSource_(
                rows=[
                    dal.ScriptSourceRow_(kind=kind, key=key, text=text)
                    for kind, key, text in definition["rows"]
                ]
            ),
            header=dal.ScriptInstrumentHeader_(
                pay_currency="USD",
                result_currency="USD",
                collateral_currency="USD",
                collateral_type="OIS",
                receiver="PAY",
                notional_parameter="N",
                settlement="AtEvent",
                valuation_cashflows="IncludeUnsettledAtValuation",
                discount_route="USD.OIS",
                convention_set="dal.fixed-dates/1",
            ),
            parameters=[
                dal.ScriptParameterDeclaration_(
                    key="N",
                    type="Real",
                    unit=unit("Price:USD"),
                    role="Notional",
                    bindable=True,
                    lower=0,
                    lower_inclusive=False,
                    upper_inclusive=True,
                ),
                dal.ScriptParameterDeclaration_(
                    key="K",
                    type="Real",
                    unit=unit(definition["unit"]),
                    role="Coupon",
                    bindable=True,
                    lower_inclusive=True,
                    upper_inclusive=True,
                ),
            ],
            measures=[
                dal.CalibrationMeasureSpec_(
                    id="Price", kind="Price", receiver="PAY", unit=unit("Price:USD")
                ),
                dal.CalibrationMeasureSpec_(
                    id=definition["measure"],
                    kind=definition["measure_kind"],
                    receiver="PAY",
                    parameter="K",
                    unit=unit(definition["unit"]),
                    annuity_sign=definition["annuity_sign"],
                    min_abs_annuity=1e-12,
                ),
            ],
        )
        return dal.ScriptCalibrationInstrument_New(
            dal.ScriptCalibrationInstrumentSpec_(
                id=definition["id"],
                definition=dal.ScriptInstrumentDefinition_New(spec),
                pillar=dal_date(definition["end"]),
            )
        )

    def observation(definition, measure=None, quote_value=None, unit_token=None):
        measure = measure or definition["measure"]
        quote = dal.CalibrationMarketQuote_New(
            definition["quote_id"],
            definition["quote"] if quote_value is None else quote_value,
            unit(unit_token or definition["unit"]),
        )
        return dal.CalibrationObservation_New(
            definition["observation_id"], instrument(definition), measure, quote, 1e-10
        )

    observations = [observation(definition) for definition in yield_instruments()]
    options = dal.CurveInstrumentCalibrationOptions_(
        evaluation="AutoFast", jacobian_mode="AAD", difference_step=1e-6
    )
    for scheme in ["LOG_LINEAR", "LOG_CUBIC_NATURAL", "MIXED"]:
        spec = dal.CurveInstrumentCalibrationSpec_(
            geometry=dal.CurveInstrumentGeometry_(
                today=dal_date(TODAY),
                currency="USD",
                curve_name="ptirds",
                knot_dates=[dal_date(value) for value in KNOTS],
                parameterization="LOG_DISCOUNT",
                knot_policy="INPUT",
                log_df_scheme=scheme,
                libor_basis="ACT_365F",
            ),
            market=dal.CurveInstrumentMarket_(
                routes=[dal.CurveScriptRoute_(id="USD.OIS")]
            ),
            valuation=dal.ScriptValuationSettings_(
                evaluation_date=dal_date(TODAY),
                fixings=dal.MarketFixingSnapshot_New({}),
            ),
            observations=observations,
            solver=dal.CurveInstrumentSolver_(
                solve_mode="EXACT",
                initial_guess=0.05,
                smoothing_weight=1.0,
                max_evaluations=200,
                max_restarts=20,
                fit_norm_tolerance=1e-8,
            ),
        )
        fit = dal.Calibrate_CurveInstruments(spec, options=options)
        assert fit.fit_within_tolerance and len(fit.parameters) == 13
        assert max(map(abs, fit.raw_residuals)) < 1e-8
    definition = xccy_instrument()
    today, end = dal_date(date(2026, 10, 21)), dal_date(date(2027, 11, 4))
    blocks = []
    for ccy, rate, tenor in [("USD", 0.04, "3M"), ("EUR", 0.02, "6M")]:
        curve = dal.DiscountPWC_New(ccy, ccy, [today, end], [rate, rate])
        blocks.append(
            dal.CurveBlock_New(
                ccy,
                ccy,
                {dal.CollateralType_OIS(): curve},
                {dal.PeriodLength_New(tenor): curve},
                dal.DayBasis_New("ACT_365F"),
            )
        )
    context = dal.XccyScriptContext_(
        valuation_time=dal.DateTime_(today, 0),
        pair=dal.CurrencyPair_New("USD", "EUR"),
        collateral_currency="USD",
        discount_route="USD.OIS",
        fixings=dal.MarketFixingSnapshot_New({}),
    )
    spec = dal.XccyInstrumentCalibrationSpec_(
        context=context,
        domestic_curve_block=blocks[0],
        foreign_curve_block=blocks[1],
        fx_spot=1.1,
        knot_dates=[end],
        observations=[observation(definition)],
    )
    fit = dal.Calibrate_XccyInstruments(
        spec,
        options=dal.CurveInstrumentCalibrationOptions_(
            evaluation="AutoFast",
            jacobian_mode="AAD",
            difference_step=1e-6,
            compute_quote_jacobian=True,
        ),
    )
    assert fit.fit_within_tolerance and abs(fit.parameters[0] - 0.003) < 1e-10
    print(
        "Validated three 13-quote yield calibrations and the XCCY basis fit through DAL."
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--validate",
        action="store_true",
        help="also execute the fixtures with the DAL Python bindings",
    )
    args = parser.parse_args()
    outputs = {
        "012.interpolate_curve_script.xlsx": generate_yield_workbook(),
        "013.xccy_script_calibration.xlsx": generate_xccy_workbook(),
    }
    for filename, workbook in outputs.items():
        destination = ROOT / "dal-excel" / "examples" / filename
        workbook.save(destination)
        print(f"Wrote {destination}")
    if args.validate:
        validate_examples()


if __name__ == "__main__":
    main()
